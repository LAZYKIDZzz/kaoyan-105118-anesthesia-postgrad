"""Extract searchable text from originals, keeping page numbers and sheet rows."""
import hashlib,json,pathlib,re,subprocess
from pypdf import PdfReader
from openpyxl import load_workbook
ROOT=pathlib.Path(__file__).resolve().parents[2]
P=ROOT/'v2-2027/research/2026-10-18';OUT=P/'pdf-text';OUT.mkdir(exist_ok=True)
rows=[]
for mf in (ROOT/'docs/admissions/2027').glob('*/manifest.json'):
    j=json.loads(mf.read_text())
    for f in j['files']:
        file=ROOT/f['path'];h=hashlib.sha256(file.read_bytes()).hexdigest();target=OUT/(j['name']+'-'+h[:8]+'.txt')
        r={'school':j['name'],**f,'text_path':str(target.relative_to(ROOT))}
        try:
            if target.exists():t=target.read_text()
            elif file.suffix=='.pdf':
                pdf=PdfReader(file);t='\n'.join('第%d页\n%s'%(i+1,p.extract_text() or '') for i,p in enumerate(pdf.pages));target.write_text(t)
            elif file.suffix in ['.xlsx','.xlsm']:
                wb=load_workbook(file,data_only=True);t='\n'.join('\n工作表：'+ws.title+'\n'+'\n'.join('行%d | '%i+' | '.join('' if c is None else str(c) for c in cells) for i,cells in enumerate(ws.iter_rows(values_only=True),1)) for ws in wb);target.write_text(t)
            elif file.suffix=='.doc':
                t=subprocess.check_output(['textutil','-convert','txt','-stdout',str(file)]).decode('utf-8');target.write_text(t)
            else:continue
            matches=list(re.finditer(r'105118|麻醉学',t));r['matches']=len(matches);r['snippets']=[t[max(0,m.start()-100):m.start()+650] for m in matches[:20]]
            rows.append(r);print(j['name'],file.name,'专业条目',len(matches),flush=True)
        except Exception as e:r['error']=str(e);rows.append(r);print('ERROR',j['name'],str(e),flush=True)
(P/'extractions.json').write_text(json.dumps(rows,ensure_ascii=False,indent=2))
